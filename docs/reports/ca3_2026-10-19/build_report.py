from pathlib import Path
import re,json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
ROOT=Path(__file__).resolve().parent
source=(ROOT/'report_source.md').read_text(encoding='utf-8-sig')
doc=Document();section=doc.sections[0]
for element in list(doc.styles.element.iter(qn('w:pBdr'))):element.getparent().remove(element)
section.page_width=Inches(8.2677);section.page_height=Inches(11.6929)
section.top_margin=section.bottom_margin=Inches(.73);section.left_margin=section.right_margin=Inches(.82)
section.footer_distance=Inches(.32)
for name,size in [('Normal',12),('Title',26),('Heading 1',17),('Heading 2',12.5),('Caption',9.5)]:
 style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
 style.paragraph_format.line_spacing=1.13;style.paragraph_format.space_after=Pt(7)
 style.paragraph_format.widow_control=True
 style._element.get_or_add_rPr().find(qn('w:rFonts')).set(qn('w:eastAsia'),'Times New Roman')
 if name.startswith('Heading'):style.paragraph_format.keep_with_next=True;style.paragraph_format.space_before=Pt(8)
doc.styles['Normal'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
doc.styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
doc.core_properties.title='EchoSpace AI Room Footprint Reconstruction'
doc.core_properties.subject='Deep Learning CA3 Project Report'
doc.core_properties.author='';doc.core_properties.keywords='EchoSpace; room reconstruction; RIR; CA3'
p=section.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run();r.font.name='Times New Roman';r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)

def table(lines):
 rows=[[c.strip() for c in line.strip().strip('|').split('|')] for line in lines]
 alignments=[('right' if ': ' in c else 'left') for c in rows[1]]
 rows=[rows[0]]+rows[2:]
 t=doc.add_table(rows=0,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=[6.62/len(rows[0])]*len(rows[0])
 if len(widths)==2:widths=[2.8,3.82]
 elif len(widths)==3:widths=[2.75,1.4,2.47]
 elif len(widths)==4:widths=[2.55,1.36,1.36,1.35]
 elif len(widths)==5:widths=[1.1,1.1,1.1,1.65,1.67]
 for col,width in zip(t.columns,widths):col.width=Inches(width)
 pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 pr.append(borders)
 for i,row in enumerate(rows):
  cells=t.add_row().cells
  trpr=cells[0]._tc.getparent().get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if i==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
  for j,(cell,text) in enumerate(zip(cells,row)):
   cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'353535' if i==0 else ('F3F5F5' if i%2==0 else 'FFFFFF'));tcpr.append(shade)
   margins=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
   tcpr.append(margins)
   p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.05
   p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or len(text)>25 else WD_ALIGN_PARAGRAPH.CENTER
   r=p.add_run(text);r.font.name='Times New Roman';r.font.size=Pt(10.5);r.bold=i==0;r.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
 spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(2);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=Pt(2);spacer.add_run().font.size=Pt(2)

pages=source.split('===PAGE===');index=[];images=[]
for number,page in enumerate(pages,1):
 # Each section starts at its own page without an extra break paragraph.
 lines=page.strip().splitlines();i=0
 while i<len(lines):
  line=lines[i].strip()
  if not line:i+=1;continue
  if line.startswith('|'):
   group=[]
   while i<len(lines) and lines[i].strip().startswith('|'):group.append(lines[i]);i+=1
   table(group);continue
  if line.startswith('!['):
   m=re.fullmatch(r'!\[(.*?)\]\((.*?)\|([\d.]+)\)',line);assert m,line
   caption,path,width=m.groups();path=(ROOT/path).resolve();width=float(width)
   p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3)
   run=p.add_run();run.add_picture(str(path),width=Inches(width));docpr=run._r.xpath('.//wp:docPr')[0];docpr.set('descr',caption)
   p=doc.add_paragraph(caption,'Caption');p.paragraph_format.keep_with_next=False;images.append(str(path));i+=1;continue
  if line.startswith('# '):
   title=line[2:];p=doc.add_paragraph(title,'Title' if number==1 else 'Heading 1');index.append({'intended_page':number,'heading':title})
   if number>1:p.paragraph_format.page_break_before=True
   if number==1:p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(36);p.paragraph_format.space_after=Pt(18)
  elif line.startswith('## '):
   title=re.sub(r'^\d+\.\d+\s+','',line[3:]);p=doc.add_paragraph(title,'Heading 2')
   if number==1:p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(30);p.runs[0].font.size=Pt(16)
  else:
   p=doc.add_paragraph(line)
   if number==1:
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i<9 else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after=Pt(15);p.paragraph_format.line_spacing=1.3
    if line.startswith('This report'):p.paragraph_format.space_before=Pt(14);p.runs[0].font.size=Pt(11)
   if number==24:
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(8)
    for r in p.runs:r.font.size=Pt(10.5)
  i+=1
out=ROOT/'EchoSpace_AI_CA3_Report.docx';doc.save(out)
(ROOT/'layout_plan.json').write_text(json.dumps({'intended_pages':len(pages),'index':index,'images':images,'words':len(source.split())},indent=2))
print(json.dumps({'docx':str(out),'planned_pages':len(pages),'images':len(images),'words':len(source.split())}))
